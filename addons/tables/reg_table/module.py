from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from engine import significance
from ui.dialogs import choose_xlsx_path
from ui.panel import Panel


# Takes as many models as are connected: a new empty input appears whenever every input is used,
# and the empty one at the end is never needed
GROWING_INPUTS = True
REQUIRED_INPUTS = 1
OUTPUTS = 1
INPUT_SPECS = [
    ("Models", "the coefficient table (Out 1) of a regression such as OLS, Logit or Fixed Effects, one per input. In 1 becomes column (1), In 2 column (2) and so on"),
]
OUTPUT_SPECS = [
    ("Table", "the table as text: a column per model with each coefficient and its stars above its standard error, then the fit statistics"),
]

# The panel's option rows: label, setting and the choices, the first being the default.
# Decimals are kept as text so they can be shown on the buttons as they are.
# Controls are listed one by one, folded into a Controls: Yes / No row (Y/N), or chosen term by
# term in their own panel (Adv)
OPTIONS = [
    ("Below each coefficient", "below", ["SE", "t"]),
    ("Brackets", "brackets", ["( )", "[ ]"]),
    ("Decimal places", "decimals", ["3", "2", "4"]),
    ("Constant", "constant", ["Show", "Hide"]),
    ("Controls (C)", "controls", ["List", "Y/N", "Adv"]),
]
SETTING_OF = {label: key for label, key, choices in OPTIONS}

CONSTANT = "Constant"
CONSTANT_NAMES = ["Intercept", "Constant", "const"]
CONTROLS_ROW = "Controls"

# Fit statistics shown under the coefficients when any model has them: label, the key in the
# coefficient table's attrs, and whether it is a count (shown whole with commas)
STATISTICS = [
    ("Observations", "observations", True),
    ("R-squared", "r_squared", False),
    ("Adjusted R-squared", "adj_r_squared", False),
    ("Within R-squared", "r_squared_within", False),
    ("Pseudo R-squared", "pseudo_r_squared", False),
    ("Entities", "entities", True),
]

# The Excel file: a serif font, a heavier rule at the top and bottom of the table and light ones
# under the column headings and above the statistics, as journals print them
FONT = "Times New Roman"
HEAVY = Side(style="medium")
LIGHT = Side(style="thin")


# Fills in the defaults. settings["names"] holds the Edit panel's renames, keyed like
# "heading 1", "dependent 1", "model 1" or "row union". settings["hidden"] holds the terms Adv hides
def set_defaults(settings):
    for label, key, choices in OPTIONS:
        settings.setdefault(key, choices[0])
    if settings["controls"] == "Yes/No":
        settings["controls"] = "Y/N"
    settings.setdefault("title", "")
    settings.setdefault("note", "")
    settings.setdefault("names", {})
    settings.setdefault("hidden", [])
    settings.setdefault("controls_row", False)


# Opens a panel with the table's options and Edit and Export buttons, previewing the table
def on_click(node):
    set_defaults(node.settings)
    panel = Panel(node, "Reg Table")
    panel.add_button("Edit", open_edit)
    panel.add_button("Export", export)
    panel.preview_title = "Table"
    panel.refresh = show_table
    show_table(panel)
    return panel


# Chooses an option. Adv also opens the panel for picking the terms shown, even if already chosen
def choose(panel, button):
    panel.node.settings[SETTING_OF[button.value]] = button.text
    if button.text == "Adv":
        panel.next_panel = terms_panel(panel.node)
    else:
        show_table(panel)


def go_back(panel, button):
    panel.next_panel = panel.node.on_click()


def open_edit(panel, button):
    panel.next_panel = edit_panel(panel.node)


# Asks where to save the Excel file and remembers it, so each run saves there again.
# Saves straight away if the workflow has already run
def export(panel, button):
    settings = panel.node.settings
    path = choose_xlsx_path()
    if not path:
        return
    settings["path"] = path
    show_table(panel)
    models = connected(panel.node.last_inputs)
    if models:
        try:
            save(build(models, settings), path)
            panel.note = f"Saved to {path}"
        except ValueError as error:
            panel.note = str(error)


# Builds the table into the panel's preview from the last run's inputs, so changes show straight
# away, and puts the note underneath. Returns the models, or an empty list if there are none
def preview_table(panel):
    panel.preview = None
    panel.note = ""
    models = connected(panel.node.last_inputs)
    if not models:
        panel.message = "Connect regression results to the inputs and run the workflow"
        return []
    try:
        table = build(models, panel.node.settings)
    except ValueError as error:
        panel.message = str(error)
        return []
    panel.preview = table
    panel.note = table.attrs["note"]
    return models


# Lists the options, the title and where the table saves to, then previews the table
def show_table(panel):
    settings = panel.node.settings
    set_defaults(settings)
    panel.clear_options()
    for label, key, choices in OPTIONS:
        panel.add_option(label, [(choice, choose, settings[key] == choice) for choice in choices])

    models = connected(panel.node.last_inputs)
    panel.lines = [
        f"Title: {settings['title'] or 'none'}",
        f"Models: {len(models) or 'none yet'}    Saves to: {settings.get('path') or 'not exported yet'}",
    ]
    preview_table(panel)


# The panel for Adv: each term any model has, with Show to put it in the table or leave it out.
# Controls row adds a Controls: Yes / No row saying which models had a term left out
def terms_panel(node):
    panel = Panel(node, "Covariates")
    panel.add_button("Back", go_back)
    panel.add_button("Show all", show_all)
    panel.add_button("Hide all", hide_all)
    panel.row_button = panel.add_button("", toggle_controls_row)
    panel.preview_title = "Table"
    panel.refresh = show_terms
    show_terms(panel)
    return panel


def controls_row_text(settings):
    return "Row: On" if settings["controls_row"] else "Row: Off"


def toggle_shown(panel, button):
    hidden = panel.node.settings["hidden"]
    if button.value in hidden:
        hidden.remove(button.value)
    else:
        hidden.append(button.value)
    show_terms(panel)


def show_all(panel, button):
    panel.node.settings["hidden"] = []
    show_terms(panel)


def hide_all(panel, button):
    panel.node.settings["hidden"] = all_terms(connected(panel.node.last_inputs))
    show_terms(panel)


def toggle_controls_row(panel, button):
    settings = panel.node.settings
    settings["controls_row"] = not settings["controls_row"]
    show_terms(panel)


# Lists every term with a Show button, as renamed in Edit, and previews the table
def show_terms(panel):
    settings = panel.node.settings
    set_defaults(settings)
    settings["controls"] = "Adv"
    panel.row_button.text = controls_row_text(settings)
    panel.clear_options()
    models = connected(panel.node.last_inputs)
    terms = all_terms(models)
    names = settings["names"]
    for term in terms:
        renamed = names.get(f"row {term}")
        panel.add_option(f"{term} ({renamed})" if renamed else term, [
            ("Show", toggle_shown, term not in settings["hidden"]),
        ])
        panel.options[-1][1][0].value = term

    shown = [term for term in terms if term not in settings["hidden"]]
    panel.lines = [
        f"Showing {len(shown)} of {len(terms)} covariates",
        f"Row: On adds a {names.get('row ' + CONTROLS_ROW) or CONTROLS_ROW} row saying which models had covariates left out",
    ]
    preview_table(panel)


# The panel for renaming: every heading, row label, the title and the note, each with an Edit
# button. The box at the top renames the one chosen, and an empty box puts back the default
def edit_panel(node):
    panel = Panel(node, "Edit Table")
    panel.add_button("Back", go_back)
    panel.add_button("Reset all", reset_names)
    panel.option_columns = 1
    panel.editing = "title"
    panel.preview_title = "Table"
    panel.refresh = show_edit
    show_edit(panel)
    return panel


# The things the Edit panel can rename, as (key, what it is, default text), for the models given
def editable(models, table):
    items = [("title", "Title", ""), ("note", "Note", note_text(table.attrs["settings"]))]
    for position, results in enumerate(models, start=1):
        items += [
            (f"heading {position}", f"Column {position} heading", f"({position})"),
            (f"dependent {position}", f"Column {position} dependent variable", dependent_of(results)),
            (f"model {position}", f"Column {position} model", model_name(results)),
        ]
    items += [(f"row {label}", "Row", label) for label in table.attrs["row_labels"]]
    return items


# The text an item has now: the title and note are kept on their own, the rest in settings["names"]
def current_text(settings, key, default):
    if key in ("title", "note"):
        return settings[key] or default
    return settings["names"].get(key, default)


def choose_item(panel, button):
    panel.editing = button.value
    panel.text_field = None
    show_edit(panel)
    panel.text_field.focused = True


# Keeps the chosen item's new text as the user types. Text matching the default, or none, is
# forgotten, so the default comes back (and follows the data if it changes)
def rename(panel, text):
    settings = panel.node.settings
    key = panel.editing
    default = dict((k, d) for k, label, d in panel.items).get(key, "")
    value = "" if text.strip() == default else text
    if key in ("title", "note"):
        settings[key] = value
    elif value.strip():
        settings["names"][key] = value
    else:
        settings["names"].pop(key, None)
    show_edit(panel)


def reset_names(panel, button):
    settings = panel.node.settings
    settings["names"] = {}
    settings["title"] = ""
    settings["note"] = ""
    panel.text_field = None
    show_edit(panel)


# Lists everything that can be renamed, showing any new name, with the box for the chosen one,
# and previews the table with the new names
def show_edit(panel):
    settings = panel.node.settings
    set_defaults(settings)
    panel.clear_options()
    panel.lines = []
    models = preview_table(panel)
    if not models:
        panel.items = []
        panel.text_field = None
        return

    panel.items = editable(models, panel.preview)
    if panel.editing not in [key for key, label, default in panel.items]:
        panel.editing = "title"
        panel.text_field = None
    for key, label, default in panel.items:
        text = current_text(settings, key, default)
        shown = f"{label}: {default or 'none'}"
        if text != default:
            shown += f" -> {text}"
        panel.add_option(shown, [("Edit", choose_item, key == panel.editing)])
        panel.options[-1][1][0].value = key

    if panel.text_field is None:
        key, label, default = next(item for item in panel.items if item[0] == panel.editing)
        name = f"Row {default}" if label == "Row" else label
        panel.add_text_field(f"{name}:", current_text(settings, key, default), rename, blocked="", fill=True)
    panel.lines = ["Click Edit on an item, then type its new name. An empty box puts back the default"]


# The coefficient tables that arrived, in input order, skipping inputs left empty
def connected(inputs):
    return [results for results in inputs or [] if results is not None]


def is_constant(row):
    return row.get("role") == "Intercept" or str(row["variable"]) in CONSTANT_NAMES


# A term's row label: its variable, with the lag for time series terms, e.g. "gdp (lag 2)"
def term_label(row):
    if is_constant(row):
        return CONSTANT
    label = str(row["variable"])
    lag = row.get("lag")
    if lag not in (None, "", 0) and not pd.isna(lag):
        label = f"{label} (lag {int(lag)})"
    return label


# Every term any model has, in the order they first appear, without the constant or period dummies
def all_terms(models):
    labels = [
        term_label(row)
        for results in models if "variable" in results.columns
        for row in results.to_dict("records") if row.get("role") != "Time"
    ]
    return [label for label in dict.fromkeys(labels) if label != CONSTANT]


# A number to a fixed number of decimal places, without a minus sign on a rounded zero
def number(value, decimals):
    if value is None or pd.isna(value):
        return ""
    text = f"{float(value):.{decimals}f}"
    return text.lstrip("-") if float(text) == 0 else text


def stars_of(row):
    if "p_value" in row and not pd.isna(row["p_value"]):
        return significance.stars(row["p_value"])
    stars = row.get("stars")
    return stars if isinstance(stars, str) else ""


# The cell under a coefficient: its standard error, or its t (or z) statistic, in brackets
def below_text(row, settings, decimals):
    if settings["below"] == "t":
        value = row.get("t_stat", row.get("z_stat"))
        if value is None and row.get("std_error"):
            value = row["coefficient"] / row["std_error"]
    else:
        value = row.get("std_error")
    text = number(value, decimals)
    if not text:
        return ""
    opening, closing = settings["brackets"][0], settings["brackets"][-1]
    return f"{opening}{text}{closing}"


# True if a term is left out of the table: the constant when hidden, controls under Y/N, and the
# terms not chosen under Adv. Period dummies from time effects are always left out, as the Time
# effects row stands for them
def is_hidden(row, label, settings):
    if row.get("role") == "Time":
        return True
    if label == CONSTANT:
        return settings["constant"] == "Hide"
    if settings["controls"] == "Y/N":
        return row.get("role") == "C"
    if settings["controls"] == "Adv":
        return label in settings["hidden"]
    return False


# One model's terms as label -> (coefficient with stars, cell below), in its table's order, and
# whether it had a term the Controls row stands for: a control under Y/N, a hidden term under Adv
def model_terms(results, settings, decimals):
    terms = {}
    left_out = False
    for row in results.to_dict("records"):
        label = term_label(row)
        if is_hidden(row, label, settings):
            left_out = left_out or (row.get("role") != "Time" and label != CONSTANT)
            continue
        top = number(row["coefficient"], decimals)
        terms[label] = (f"{top}{stars_of(row)}" if top else "", below_text(row, settings, decimals))
    return terms, left_out


def model_name(results):
    return results.attrs.get("model") or results.attrs.get("fit", {}).get("model", "")


def dependent_of(results):
    return results.attrs.get("fit", {}).get("dependent") or ""


def errors_name(fit):
    cluster = fit.get("cluster")
    if cluster:
        return f"Clustered ({cluster['column']})"
    return "Robust" if fit.get("robust") else "Ordinary"


# The rows under the coefficients: each statistic any model has, then the controls, panel
# effects and kind of standard errors where they tell the models apart from plain OLS
def statistic_rows(models, left_out, settings, decimals):
    rows = []
    for label, key, count in STATISTICS:
        values = [results.attrs.get(key) for results in models]
        if all(value is None for value in values):
            continue
        cells = [
            "" if value is None or pd.isna(value) else f"{int(value):,}" if count else number(value, decimals)
            for value in values
        ]
        rows.append([label] + cells)

    controls_row = settings["controls"] == "Y/N" or (settings["controls"] == "Adv" and settings["controls_row"])
    if controls_row and any(left_out):
        rows.append([CONTROLS_ROW] + ["Yes" if hidden else "No" for hidden in left_out])

    if any("entities" in results.attrs for results in models):
        effects = {"Fixed Effects": "Fixed", "Random Effects": "Random"}
        rows.append(["Entity effects"] + [effects.get(model_name(results), "No") for results in models])
        rows.append(["Time effects"] + ["Yes" if results.attrs.get("time_effects") else "No" for results in models])

    errors = [errors_name(results.attrs) for results in models]
    if any(name != "Ordinary" for name in errors):
        rows.append(["Standard errors"] + errors)
    return rows


# What the brackets hold and what the stars mean, most stars first, as journals put it
def note_text(settings):
    below = "Standard errors" if settings["below"] == "SE" else "t statistics"
    brackets = "parentheses" if settings["brackets"] == "( )" else "square brackets"
    levels = list(enumerate(significance.levels, start=1))[::-1]
    stars = ", ".join(f"{'*' * count} p < {level:g}" for count, level in levels)
    return f"Notes: {below} in {brackets}. {stars}."


# Builds the table: a label column then one column per model, numbered (1), (2) and so on.
# The first rows give each model's dependent variable and name, then each term gets a row with
# its coefficients and one under it in brackets, with the constant last, then the statistics.
# Headings and row labels take any new names from Edit. Every cell is text so the stars and
# brackets stay as they are. The attrs say how many rows each block has, so the Excel file can
# rule them off, plus the note, title and the row labels before renaming
def build(models, settings):
    set_defaults(settings)
    for position, results in enumerate(models, start=1):
        missing = [column for column in ["variable", "coefficient"] if column not in results.columns]
        if missing:
            raise ValueError(f"Model {position} is not a coefficient table: it has no {' or '.join(missing)} column")

    names = settings["names"]
    positions = range(1, len(models) + 1)
    decimals = int(settings["decimals"])
    terms, left_out = zip(*[model_terms(results, settings, decimals) for results in models])
    labels = list(dict.fromkeys(label for model in terms for label in model if label != CONSTANT))
    if any(CONSTANT in model for model in terms):
        labels.append(CONSTANT)

    header = []
    dependents = [names.get(f"dependent {p}", dependent_of(results)) for p, results in zip(positions, models)]
    if any(dependents):
        header.append([""] + dependents)
    models_row = [names.get(f"model {p}", model_name(results)) for p, results in zip(positions, models)]
    if any(models_row):
        header.append([""] + models_row)

    body = []
    for label in labels:
        cells = [model.get(label, ("", "")) for model in terms]
        body.append([label] + [top for top, below in cells])
        body.append([""] + [below for top, below in cells])

    statistics = statistic_rows(models, left_out, settings, decimals)
    row_labels = [row[0] for row in body + statistics if row[0]]
    for row in body + statistics:
        if row[0]:
            row[0] = names.get(f"row {row[0]}", row[0])

    columns = [""] + [names.get(f"heading {p}", f"({p})") for p in positions]
    table = pd.DataFrame(header + body + statistics, columns=columns)
    table.attrs = {
        "header_rows": len(header),
        "coefficient_rows": len(body),
        "statistic_rows": len(statistics),
        "note": settings["note"] or note_text(settings),
        "title": settings["title"],
        "row_labels": row_labels,
        "settings": {key: settings[key] for key in ["below", "brackets"]},
    }
    return table


# Writes the table to an Excel file laid out as a journal prints it: the title, the column
# headings, the coefficients and the statistics, ruled off from each other, and the note underneath
def save(table, path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Table"
    width = len(table.columns)
    font = Font(name=FONT, size=11)
    row = 1

    title = table.attrs.get("title")
    if title:
        cell = sheet.cell(row=row, column=1, value=title)
        cell.font = Font(name=FONT, size=12, bold=True)
        cell.alignment = Alignment(horizontal="center")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        row += 2

    top = row
    lines = [list(table.columns)] + table.values.tolist()
    for offset, line in enumerate(lines):
        for column, value in enumerate(line, start=1):
            cell = sheet.cell(row=top + offset, column=column, value=value)
            cell.font = font
            cell.number_format = "@"
            cell.alignment = Alignment(horizontal="left" if column == 1 else "center")

    # Rules: heavy above the headings and under the last row, light under the headings and
    # above the statistics
    headings_end = top + table.attrs["header_rows"]
    coefficients_end = headings_end + table.attrs["coefficient_rows"]
    bottom = top + len(lines) - 1
    rules = {top: {"top": HEAVY}}
    rules.setdefault(headings_end, {})["bottom"] = LIGHT
    if table.attrs["statistic_rows"]:
        rules.setdefault(coefficients_end, {})["bottom"] = LIGHT
    rules.setdefault(bottom, {})["bottom"] = HEAVY
    for line, sides in rules.items():
        for column in range(1, width + 1):
            sheet.cell(row=line, column=column).border = Border(**sides)

    note = sheet.cell(row=bottom + 1, column=1, value=table.attrs["note"])
    note.font = Font(name=FONT, size=9)
    note.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    sheet.merge_cells(start_row=bottom + 1, start_column=1, end_row=bottom + 1, end_column=width)

    label_width = max(len(str(value)) for value in table.iloc[:, 0].tolist() + [""])
    sheet.column_dimensions["A"].width = max(14, label_width + 2)
    for column in range(2, width + 1):
        cells = [str(value) for value in table.iloc[:, column - 1].tolist() + [table.columns[column - 1]]]
        letter = sheet.cell(row=top, column=column).column_letter
        sheet.column_dimensions[letter].width = max(12, max(len(cell) for cell in cells) + 3)
    total = sum(sheet.column_dimensions[sheet.cell(row=top, column=c).column_letter].width for c in range(1, width + 1))
    sheet.row_dimensions[bottom + 1].height = 13 * (len(table.attrs["note"]) // max(1, int(total * 1.1)) + 1)
    sheet.sheet_view.showGridLines = False

    try:
        book.save(path)
    except PermissionError:
        raise ValueError(f"Could not save {path}: close it if it is open in Excel, then try again")


# Builds the table from the connected coefficient tables, and saves it to Excel if a file was chosen
def run(inputs, settings):
    table = build(connected(inputs), settings)
    path = settings.get("path")
    if path:
        if not Path(path).parent.exists():
            raise ValueError(f"The folder for {path} no longer exists: click Export to choose another")
        save(table, path)
        print(f"Saved the regression table to {path}")
    return [table]
